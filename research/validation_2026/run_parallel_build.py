"""Parallel rebuild of the reference database from staged EDFs.

Splits the file set across N workers (each builds its own partial out_file with
build_db_from_edfs.py), waits, then merges all per-file columns into one
database and recomputes AVG / STD DEV with the same trimmed estimator the
pipeline uses (drop top/bottom 5%, mean / population std).

Usage: py run_parallel_build.py <src_folder> <database_name> [n_workers]
"""
import math
import os
import subprocess
import sys
import time

import numpy as np
import openpyxl

HERE = os.path.dirname(__file__)
REPO_ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
OUT = os.path.join(HERE, "out")
DRIVER = os.path.join(HERE, "build_db_from_edfs.py")

import re
VARIANT = re.compile(r"\.(clean1?|recon|unclean)\.edf$", re.I)
ECRE = re.compile(r"(?<![A-Za-z])EC(?![A-Za-z])", re.I)


def count_edfs(folder):
    # Count every non-variant EDF (staging already selected the condition).
    return len([f for f in os.listdir(folder)
                if f.lower().endswith(".edf") and not VARIANT.search(f)
                and "disregard" not in f.lower()])


def trimmed(vals):
    v = np.sort(np.asarray([x for x in vals if x is not None and x == x], float))
    n = len(v)
    k = int(0.05 * n)
    core = v[k:n - k] if k > 0 and n - k > k else v
    if core.size == 0:
        core = v
    return float(core.mean()), float(core.std())


def out_path(dbname):
    return os.path.join(REPO_ROOT, dbname + ".out_file.icale.xlsx")


def main():
    src = sys.argv[1]
    dbname = sys.argv[2]
    nworkers = int(sys.argv[3]) if len(sys.argv) > 3 else 8

    total = count_edfs(src)
    chunk = math.ceil(total / nworkers)
    print(f"[parallel] {total} files, {nworkers} workers, chunk={chunk}", flush=True)

    env = os.environ.copy()
    for k in ("OMP_NUM_THREADS", "MKL_NUM_THREADS", "OPENBLAS_NUM_THREADS"):
        env[k] = "2"  # limit BLAS threads per worker to avoid oversubscription

    procs = []
    t0 = time.time()
    for w in range(nworkers):
        start = w * chunk
        if start >= total:
            break
        part = f"{dbname}_p{w}"
        for suf in (".out_file.icale.xlsx", ".names_file.icale.xlsx"):
            try:
                os.remove(os.path.join(REPO_ROOT, part + suf))
            except OSError:
                pass
        log = open(os.path.join(OUT, f"_pbuild_p{w}.log"), "w")
        p = subprocess.Popen(
            [sys.executable, DRIVER, src, part, str(start), str(chunk)],
            stdout=log, stderr=subprocess.STDOUT, env=env, cwd=REPO_ROOT)
        procs.append((w, part, p, log))
        print(f"[parallel] launched worker {w}: files [{start}:{start+chunk}]", flush=True)

    fails = 0
    for w, part, p, log in procs:
        rc = p.wait()
        log.close()
        print(f"[parallel] worker {w} done rc={rc} ({(time.time()-t0)/60:.1f} min)", flush=True)
        if rc != 0:
            fails += 1

    # ---- merge partials
    print("[parallel] merging partials ...", flush=True)
    template = None
    data_cols = None  # list of rows, each a list of per-file values
    total_cols = 0
    for w, part, p, log in procs:
        pth = out_path(part)
        if not os.path.exists(pth):
            print(f"[parallel] WARN missing partial {part}", flush=True)
            continue
        rows = list(openpyxl.load_workbook(pth, data_only=True, read_only=True).active
                    .iter_rows(values_only=True))
        if template is None:
            template = rows
            data_cols = [[] for _ in rows]
        for ri, row in enumerate(rows):
            data_cols[ri].extend(row[2:])  # skip AVG/STD columns
        total_cols += len(rows[0]) - 2

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["AVG", "STD DEV"] + [f"C{i+1}" for i in range(total_cols)])
    for ri in range(1, len(template)):        # row 0 is the header
        vals = data_cols[ri]
        avg, std = trimmed(vals)
        ws.append([round(avg, 6), round(std, 6)] +
                  [round(float(x), 6) if isinstance(x, (int, float)) else None for x in vals])
    wb.save(out_path(dbname))
    print(f"[parallel] merged {total_cols} files -> {out_path(dbname)} "
          f"({fails} worker failures) in {(time.time()-t0)/60:.1f} min", flush=True)


if __name__ == "__main__":
    main()
