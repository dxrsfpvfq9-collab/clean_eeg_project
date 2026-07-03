"""Build the study Comparison Chart (Brain Panel vs Doctor's Report), matching
the layout of the original first-study spreadsheet
('04_Comparison Chart BP to DR ... 07AUG.xlsx').

Columns:
  A Patient ID  B Client ID
  C BP Total OOB (=SUM(D:I))  D Std/Global(2)  E PDR(9)  F Phenotype(6)
  G Focal/Frontal(12)  H Diffuse(7)  I State Shift(12)
  J PDR Magnitude use Alpha Peak?   K BP Comments (machine findings)
  L Overall quality  M Presence of artifact  N Amount artifact-free
  O Background rhythm  P Drowsiness/Sleep  Q Paroxysmal disturbance
  R DR Comments  S Doctor full text

One row per cohort EC panel. OOB counts come from the saved Brain Panel PDF;
L-S from the paired Doctor's Report .docx.
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
import panel_parser as pp
import report_parser as rp

sys.path.insert(0, config.REPO_ROOT)
from process import discriminant  # noqa: E402

OUT_XLSX = os.path.join(config.OUT_DIR, "comparison_chart_new_data.xlsx")

HEADERS = {
    "C": "BP Total # Metrics out of bounds -  Total 48",
    "D": "STD/Global Raw-2 total",
    "E": "PDR Metrics out of Bounds - 9 total",
    "F": "Phenotype out of bounds - 6 total",
    "G": "Focal/Frontal Out of Bounds - 12 total ",
    "H": "Diffuse Out of bounds - 7 total ",
    "I": "State Shift out of Bounds - 12 total",
    "J": "PDR Magnitude use Alpha Peak?",
    "K": "BP Comments",
    "L": "DR Comments Overall quality",
    "M": "Presence of artifact",
    "N": "Amount of artifact free portions",
    "O": "Background rhythym",
    "P": "Drowsiness/Sleep",
    "Q": "Paroxsmal Disturbance",
    "R": "DR Comments",
    "S": "Doc full Text",
}
SUBNOTES = {  # row-2 "Corresponds to" notes from the original
    "C": "Corresponds to: Overall quality, inadequate due to artifact",
    "D": "Corresponds to: Overall quality, inadequate due to artifact",
    "E": "Corresponds to background rhythym",
    "G": "Corresponds to Background Rhythym ",
    "H": "Corresponds to artifact",
    "I": "Corresponds to drowsiness",
    "J": "Corresponds to PDR as reported by DR",
    "K": "Corresponds to DR Comments",
}
WIDE = {"K": 55, "M": 40, "O": 55, "P": 40, "R": 45, "S": 70, "L": 14, "N": 22, "Q": 18}


def client_id(row):
    fn = row["ec_panel_file"]
    return fn.split(".icale.rep.pdf")[0].strip()


def main():
    rows = list(csv.DictReader(open(os.path.join(config.OUT_DIR, "study_cohort.csv"),
                                    encoding="utf-8")))
    print(f"Building comparison chart for {len(rows)} cohort panels ...")

    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison Chart"
    arial = "Arial"

    hdr_font = Font(name=arial, bold=True, size=10)
    sub_font = Font(name=arial, italic=True, size=8, color="666666")
    cell_font = Font(name=arial, size=10)
    wrap_top = Alignment(wrap_text=True, vertical="top")
    ctr = Alignment(horizontal="center", vertical="top")
    fill = PatternFill("solid", fgColor="DDEBF7")

    # row 1 headers
    for col, txt in HEADERS.items():
        c = ws[f"{col}1"]; c.value = txt; c.font = hdr_font
        c.alignment = wrap_top; c.fill = fill
    # row 2 subnotes
    for col, txt in SUBNOTES.items():
        c = ws[f"{col}2"]; c.value = txt; c.font = sub_font; c.alignment = wrap_top
    # row 3 A/B labels
    for col, txt in (("A", "Patient ID"), ("B", "Client ID")):
        c = ws[f"{col}3"]; c.value = txt; c.font = hdr_font; c.fill = fill

    n_template_ok = 0
    r = 4
    for i, row in enumerate(rows, 1):
        parsed = pp.parse_panel_pdf(row["ec_panel_path"])
        counts = discriminant.compute_oob_counts(parsed["zscores"])
        findings = pp.extract_findings(row["ec_panel_path"])
        rep = rp.parse_report(row["report_path"]) if row["report_path"] else \
            {k: "" for k in rp.FIELD_KEYS + ["full_text"]}
        if rep.get("template_ok"):
            n_template_ok += 1

        ws[f"A{r}"] = f"BP{i:03d}"
        ws[f"B{r}"] = client_id(row)
        ws[f"C{r}"] = f"=SUM(D{r}:I{r})"            # total = sum of groups
        ws[f"D{r}"] = counts["std_global"]
        ws[f"E{r}"] = counts["pdr"]
        ws[f"F{r}"] = counts["phenotypes"]
        ws[f"G{r}"] = counts["focal"]
        ws[f"H{r}"] = counts["diffuse"]
        ws[f"I{r}"] = counts["state_shift"]
        ws[f"J{r}"] = ""                             # manual reviewer field
        ws[f"K{r}"] = findings
        ws[f"L{r}"] = rep.get("quality", "")
        ws[f"M{r}"] = rep.get("artifact", "")
        ws[f"N{r}"] = rep.get("artifact_free", "")
        ws[f"O{r}"] = rep.get("background", "")
        ws[f"P{r}"] = rep.get("drowsiness", "")
        ws[f"Q{r}"] = rep.get("paroxysmal", "")
        ws[f"R{r}"] = rep.get("comments", "")
        ws[f"S{r}"] = rep.get("full_text", "")
        for col in [get_column_letter(x) for x in range(1, 20)]:
            cell = ws[f"{col}{r}"]
            cell.font = cell_font
            cell.alignment = ctr if col in "CDEFGHI" else wrap_top
        r += 1

    # widths + freeze
    for col in [get_column_letter(x) for x in range(1, 20)]:
        ws.column_dimensions[col].width = WIDE.get(col, 11)
    ws.freeze_panes = "C4"

    wb.save(OUT_XLSX)
    print(f"  wrote {OUT_XLSX}")
    print(f"  rows: {len(rows)} | reports parsed (template_ok): {n_template_ok}")


if __name__ == "__main__":
    main()
