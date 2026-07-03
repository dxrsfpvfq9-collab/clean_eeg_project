# -*- coding: utf-8 -*-
"""Build the Clinical chapter (Ims, Collura, Turner) per Rusty's email split."""
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"C:/Users/tcollura/Dropbox/Documents/Papers/Epilepsy Book Tato FNNR/Ims_Collura_Turner_Epileptiform_Activity_NonEpileptics.docx"

# ---- palette (matches the companion treatment chapter) ----
C_TITLE   = RGBColor(0x1F, 0x3A, 0x5F)
C_GREY    = RGBColor(0x5B, 0x6B, 0x7B)
C_DARK    = RGBColor(0x1A, 0x1A, 0x1A)
C_H2      = RGBColor(0x2E, 0x86, 0xAB)
RULE      = "2E86AB"

doc = Document()

# default font
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
rpr = normal.element.get_or_add_rPr().get_or_add_rFonts()
rpr.set(qn("w:ascii"), "Calibri"); rpr.set(qn("w:hAnsi"), "Calibri")

sec = doc.sections[0]
sec.page_width  = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)

def run(p, text, *, bold=False, italic=False, color=None, size=None):
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    if color is not None: r.font.color.rgb = color
    if size is not None: r.font.size = Pt(size)
    return r

def center(spacing_after=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if spacing_after is not None:
        p.paragraph_format.space_after = Pt(spacing_after)
    return p

def divider():
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), RULE)
    pbdr.append(bottom)
    # pBdr must precede w:spacing in the schema order
    spacing = pPr.find(qn("w:spacing"))
    if spacing is not None:
        spacing.addprevious(pbdr)
    else:
        pPr.insert(0, pbdr)

def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    _set_outline(p, 0)
    run(p, text, bold=True, color=C_TITLE, size=16)

def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    _set_outline(p, 1)
    run(p, text, bold=True, color=C_H2, size=13)

def _set_outline(p, lvl):
    pPr = p._p.get_or_add_pPr()
    o = OxmlElement("w:outlineLvl"); o.set(qn("w:val"), str(lvl)); pPr.append(o)

def body(text):
    """Justified body paragraph. *word* -> italic span."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    parts = text.split("*")
    for i, seg in enumerate(parts):
        if seg == "": continue
        run(p, seg, italic=(i % 2 == 1))
    return p

def page_number_footer():
    ftr = sec.footer
    p = ftr.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); r.font.color.rgb = C_GREY; r.font.size = Pt(9)
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    r._r.append(fld1); r._r.append(instr); r._r.append(fld2)

# ============================== TITLE BLOCK ==============================
p = center(4); run(p, "EEG Epileptiform Activity in People without Epilepsy: Clinical Management of Abnormal Paroxysmal Activity in Routine EEG", bold=True, color=C_TITLE, size=18)
p = center(10); run(p, "Detection and Clinical Implications of Isolated Epileptiform Discharges (IEDs) in Resting-State EEG", italic=True, color=C_GREY, size=12.5)
p = center(2); run(p, "David Ims, M.A., LCPC, BCN, QEEG-DL¹   ·   Thomas F. Collura, Ph.D., P.E., QEEG-D, BCN²   ·   Robert P. Turner, M.D., M.S.C.R., QEEG-D³", bold=True, color=C_DARK, size=11.5)
p = center(2)
run(p, "¹ Chesapeake Neuro-Behavioral Health, LLC, USA", italic=True, color=C_GREY, size=9.5); p.add_run().add_break()
run(p, "² BrainMaster Technologies, Inc. / Brain Enrichment Center, Bedford, Ohio, USA", italic=True, color=C_GREY, size=9.5); p.add_run().add_break()
run(p, "³ Network Neurology / Network Neuroscience, Charleston, South Carolina, USA", italic=True, color=C_GREY, size=9.5)
p = center(10); run(p, "Chapter prepared for Advances in Epilepsy: From Diagnosis to Treatment (E. Sokhadze, R. Turner, & M. B. Dellinger, Eds.). Foundation for Neurofeedback and Neuromodulation Research (FNNR).", italic=True, color=C_GREY, size=9)
divider()

# ============================== ABSTRACT ==============================
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after = Pt(3)
run(p, "Abstract", bold=True, color=C_TITLE, size=12.5)
body("The electroencephalogram (EEG) is the central laboratory tool in the evaluation of epilepsy, but its findings do not belong to people with epilepsy alone. Sharply contoured, paroxysmal waveforms that resemble interictal epileptiform discharges (IEDs) are encountered, with non-trivial frequency, in the resting-state EEG of individuals who have never had a seizure — and the growth of quantitative EEG (QEEG) and neurofeedback has brought many such recordings under the eye of clinicians who are not primarily epileptologists. This chapter addresses the detection and clinical implications of epileptiform and abnormal paroxysmal activity in persons without epilepsy. We first set the routine and quantitative EEG in context: how each is used in the diagnosis and classification of epilepsy, and what the interictal record does and does not establish. We then define epileptiform activity and the isolated epileptiform discharge, distinguish genuine discharges from the benign variants and normal patterns that mimic them, summarize the reported prevalence of such findings across healthy and clinical populations, and examine the groups — children, neurodevelopmental, and psychiatric populations — in which they most often arise. We give particular attention to the clinical costs of over-reading, to the defensible management of an incidental epileptiform finding in routine practice, and to what such activity means for assessment, protocol selection, and safety in QEEG-guided neurofeedback — the bridge into the companion treatment chapter of this volume.")
divider()

# ============================== 1. INTRODUCTION ==============================
h1("1.  Introduction: Paroxysmal Activity Without Epilepsy")
body("The routine EEG occupies a peculiar position in clinical neuroscience. It is the test most closely identified with epilepsy, yet its relationship to the diagnosis is profoundly asymmetric. A clearly ictal recording can confirm epilepsy, but a normal interictal study can never exclude it, and — the concern of this chapter — the appearance of a sharp, paroxysmal waveform in the record of a person who has never had a seizure does not, by itself, establish that they have epilepsy or ever will. Epilepsy is a clinical diagnosis, made from the history of recurrent unprovoked seizures and supported, not supplanted, by the EEG (Smith, 2005; Pillai & Sperling, 2006). The waveform is evidence; it is not a verdict.")
body("This distinction has acquired new practical weight. The same instrumentation that records the clinical EEG now drives a large and growing volume of quantitative EEG (QEEG) and neurofeedback assessment, much of it acquired and reviewed by practitioners whose primary training is not in epileptology. A screening QEEG obtained for attention, mood, or performance concerns is a resting-state EEG, and it will occasionally contain transients that are sharply contoured, focal, and paroxysmal — in short, that look epileptiform. How these are recognized, named, and acted upon has consequences that reach well beyond the recording room: a careless reading can label a healthy person as epileptic, with cascading effects on medication, driving, employment, insurance, and self-concept.")
body("Our purpose is therefore twofold. First, to equip the reader to *detect* and correctly characterize epileptiform and abnormal paroxysmal activity in the resting EEG — including the crucial skill of separating genuine epileptiform discharges from the many benign variants that imitate them. Second, to lay out the *clinical management* of such findings: what an isolated epileptiform discharge in a non-epileptic person does and does not predict, how it should be documented and communicated, and how it bears on the conduct of QEEG-guided neurofeedback. Throughout, the governing principle is conservative: in a person without a clinical history of seizures, the threshold for calling a waveform epileptiform — and the threshold for letting that call change the person's life — should both be high.")
divider()

# ============================== 2. CLINICAL EEG/QEEG IN CONTEXT ==============================
h1("2.  The Clinical EEG and QEEG in Context")
body("Before epileptiform activity in the well can be interpreted, the role of the EEG in the sick must be clear. This section establishes that diagnostic backdrop: how the routine and quantitative EEG are used in epilepsy, what the interictal and ictal records reveal, and where quantitative methods help and where they mislead.")

h2("2.1  The EEG in the diagnosis and classification of epilepsy")
body("The routine scalp EEG samples roughly twenty to thirty minutes of cortical activity through the standard 10–20 electrode array, ordinarily with periods of hyperventilation and photic stimulation. Its great strength is temporal resolution; its great limitation is that interictal epileptiform discharges are intermittent, so a single routine study captures them in only a minority of people who genuinely have epilepsy — sensitivities on the order of one-quarter to one-half are typical, rising substantially with repeated studies, with sleep, and with sleep deprivation (Smith, 2005; Pillai & Sperling, 2006). Ambulatory and long-term video-EEG monitoring extend the sampling window and, by capturing habitual events, allow the electrographic and clinical features of a seizure to be correlated directly. The corollary that anchors this whole chapter follows immediately: because the interictal EEG is so insensitive, clinicians learn to treat a normal study as uninformative rather than reassuring — and, by the same logic, must resist treating an isolated abnormal-looking transient as diagnostic.")
body("The International League Against Epilepsy (ILAE) provides the framework within which EEG findings are placed. The 2017 operational classification organizes seizures by their site of onset (focal, generalized, or unknown) and the epilepsies by type and, where possible, by syndrome and etiology (Fisher et al., 2017; Scheffer et al., 2017). The EEG contributes to this scheme — a generalized spike-and-wave pattern supports a genetic generalized epilepsy, a focal discharge supports a focal epilepsy — but it contributes as one line of evidence integrated with semiology, imaging, and course, never as a stand-alone classifier.")

h2("2.2  Interictal and ictal electrographic signatures")
body("The interictal hallmark of the epileptic cortex is the epileptiform discharge: the spike (a pointed transient under 70 ms), the sharp wave (70–200 ms), and their combination with a following slow wave as the spike-and-wave complex. These may be focal, reflecting a circumscribed irritative zone, or generalized and bisynchronous, as in the 3-Hz spike-and-wave of absence epilepsy. The ictal record — the electrographic seizure itself — shows an evolving, rhythmic discharge that changes in frequency, amplitude, and spatial distribution over seconds to minutes, a temporal evolution that is itself a key criterion separating a true seizure from a brief interictal burst.")
body("Background features carry their own information. Focal slowing in the delta or theta range suggests localized cerebral dysfunction and can have localizing value; diffuse slowing points to a generalized encephalopathic process. These background abnormalities matter to the present topic because they, too, are read in context: slowing in an alert, asymptomatic adult demands a different interpretation than the same finding in an obtunded patient.")

h2("2.3  Quantitative EEG (QEEG) in epilepsy")
body("Quantitative EEG applies digital signal processing to the recorded waveform — spectral (Fourier) analysis of band power, measures of inter-regional coherence and phase, and indices of hemispheric asymmetry — and compares the result against age-referenced normative databases, expressing each measure as a Z-score, the number of standard deviations by which the individual departs from expectation (Thatcher, 1998; Nuwer, 1997). In epilepsy, QEEG can quantify interictal slowing and connectivity disturbance that extend beyond a visually obvious focus, and it furnishes the individualized deviation maps that guide the neurofeedback methods of the companion chapter.")
body("Its limits must be stated as plainly as its strengths, because they bear directly on the over-reading problem. Position statements from the American Academy of Neurology and the American Clinical Neurophysiology Society caution that QEEG is an adjunct to, not a replacement for, expert visual analysis, and specifically that the *detection of epileptiform transients remains a task for the trained human reader* rather than for spectral software (Nuwer, 1997). Frequency-domain summaries average over time and are blind to the brief, sharply contoured events that define epileptiform activity; an automated spectral display can look entirely unremarkable over a segment that contains a clear spike, and conversely can flag deviations that have no epileptogenic meaning. QEEG informs; it does not adjudicate.")

h2("2.4  Source localization and network perspectives")
body("Source-estimation methods such as low-resolution electromagnetic tomography (LORETA) project scalp signals onto a model of cortical generators, moving the description from electrode sites toward estimated brain regions and networks (Pascual-Marqui, Michel, & Lehmann, 1994). Coupled with connectivity analysis, they support the modern view of the epilepsies as network disorders rather than purely focal ones, and where available, magnetoencephalography (MEG) and EEG-fMRI add complementary spatial detail. For the purposes of this chapter, the network perspective is a caution as much as a tool: a single estimated source for one sharp transient, in a person without seizures, should not be over-interpreted as evidence of an epileptogenic network.")
divider()

# ============================== 3. DEFINING ==============================
h1("3.  Defining Epileptiform Activity and Isolated Epileptiform Discharges (IEDs)")
body("“Epileptiform” is a morphological term, not a diagnosis — it describes how a waveform looks, and the word deliberately stops short of asserting what it means clinically. Standardized terminology developed for clinical electroencephalographers, most recently the revised glossary endorsed through the International Federation of Clinical Neurophysiology, defines an epileptiform discharge by a convergent set of features rather than by any single one (Kane et al., 2017). A genuine epileptiform discharge is typically di- or triphasic with a sharply pointed peak; it stands out from and disrupts the ongoing background; it has a duration distinct from that background; it is commonly followed by a slow after-going wave; and it has a physiologically plausible electrical field across more than one electrode. A pointed transient meeting few of these criteria — a monophasic blip seen at a single electrode with no field and no after-going slow wave — should not be called epileptiform.")
body("The *isolated epileptiform discharge* (IED) is, then, such a discharge occurring on its own, outside any clinical or electrographic seizure, in an interictal record. When IEDs appear in the EEG of a person with a compatible seizure history, they support the diagnosis and aid classification. When they appear in a person *without* such a history — the situation that defines this chapter — their meaning is far more guarded, and the first analytic task is to be sure they are epileptiform at all.")
body("That task is dominated by the benign variants and normal patterns of uncertain significance: sharply contoured, sometimes rhythmic patterns that mimic epileptiform activity but carry no association with seizures (Tatum, Husain, Benbadis, & Kaplan, 2006). The most important to recognize are wicket spikes (arciform temporal waveforms of drowsiness, frequently misread as temporal sharp waves); benign epileptiform transients of sleep (BETS), also called small sharp spikes; rhythmic mid-temporal theta of drowsiness (the “psychomotor variant”); 14-and-6-Hz positive bursts; the 6-Hz “phantom” spike-and-wave; and the subclinical rhythmic electrographic discharge of adults (SREDA). Each has a characteristic morphology, state-dependence, and distribution, and each lacks the disruptive after-going slow wave and evolving field of a true discharge. Familiarity with this catalogue is the single most protective skill in reading the EEG of an asymptomatic person, because nearly every false-positive diagnosis of epilepsy traces back to one of these patterns being mistaken for the real thing.")
divider()

# ============================== 4. PREVALENCE ==============================
h1("4.  Prevalence in Persons Without Epilepsy")
body("Epileptiform discharges are not confined to people with epilepsy, and the literature quantifying their occurrence in others spans more than half a century. The classic reference is Zivin and Ajmone Marsan (1968), who reviewed the EEGs of non-epileptic patients and found unequivocally epileptiform activity in roughly two percent, the large majority of whom did not go on to develop seizures over follow-up. Subsequent series in genuinely healthy populations report rates that are lower still but non-zero: in screened candidates for aircrew training — young, healthy adults with strong incentives to appear well — epileptiform abnormalities were found in well under one percent (Gregory, Oates, & Merry, 1993).")
body("Children are the conspicuous exception, showing markedly higher rates than adults. In a longitudinal study of normal children, Cavazzuti, Cappella, and Nalin (1980) detected epileptiform EEG patterns in several percent of healthy schoolchildren, and — the critical finding — most of these children never developed epilepsy, the discharges frequently disappearing with maturation. Community- and clinic-based analyses reinforce the central message that the prognostic weight of an incidentally discovered discharge in an asymptomatic person is modest (Sam & So, 2001; So, 2010).")
body("Two methodological points govern any such number. First, the measured prevalence depends heavily on technique and on the reader: montage, the inclusion of sleep and sleep deprivation, recording duration, and — above all — the criteria and conservatism of the interpreter can move the apparent rate several-fold, and much of the variation across studies is variation in reading, not in biology. Second, and following directly, the rates that matter clinically are those generated by readers applying strict, standardized morphological criteria; permissive reading inflates prevalence precisely by absorbing the benign variants of Section 3 into the epileptiform category.")
divider()

# ============================== 5. POPULATIONS ==============================
h1("5.  Populations and Clinical Contexts")
body("The likelihood and meaning of epileptiform activity without epilepsy vary by population. In children, as noted, both the prevalence and the rate of spontaneous resolution are higher, reflecting the maturing brain's changing excitability; an incidental discharge in a developing child warrants correlation with development and behavior rather than reflexive alarm.")
body("Neurodevelopmental conditions, autism spectrum disorder (ASD) in particular, show elevated rates of epileptiform EEG abnormality even in the absence of clinical seizures, and the clinical significance of these subclinical discharges remains genuinely unsettled — whether they contribute to the behavioral phenotype, merely co-occur with it, or warrant treatment when seizures are absent is an open and actively debated question (Spence & Schneider, 2009). Psychiatric and neurobehavioral populations likewise display higher rates of nonspecific dysrhythmic and occasionally sharp activity than the general population, a long-observed association whose interpretation is complicated by medication effects, state, and referral bias (Shelley, Trimble, & Boutros, 2008).")
body("Across all of these contexts the predictive content of an isolated discharge is limited. In a truly asymptomatic adult, the rate of subsequent epilepsy following an incidentally discovered discharge is low; the discharge is a statistical risk marker at most, not a diagnosis-in-waiting (Sam & So, 2001; So, 2010). The exceptions — contexts in which a specific discharge pattern carries real predictive weight — are defined by syndrome and clinical setting (for example, certain patterns after a first unprovoked seizure, where IEDs do raise recurrence risk), and they are exactly the situations in which an epileptologist, not a screening QEEG, should be guiding interpretation.")
divider()

# ============================== 6. OVER-READING ==============================
h1("6.  Interpretation, Over-Reading, and Pitfalls")
body("The dominant error in this domain is not failing to find epileptiform activity but finding too much of it. Over-reading of the EEG — most often the misidentification of a benign variant or a fragment of normal background as an epileptiform discharge — is a leading and well-documented cause of the misdiagnosis of epilepsy, and a substantial fraction of patients referred to specialist centers for “refractory epilepsy” prove, on expert review, not to have epilepsy at all (Benbadis & Tatum, 2003; Benbadis, 2007). The consequences are not abstract: unnecessary antiseizure medication with its side effects, loss of driving privileges and employment, restrictions on activity, raised insurance costs, the psychological burden of a serious chronic-disease label, and the foreclosing of the search for the patient's actual problem.")
body("The asymmetry of error is the heart of the matter. Missing a discharge in an asymptomatic person usually costs little, because the discharge in isolation rarely changes management; manufacturing one can cost a person their license, their job, and years on a drug they never needed. This asymmetry argues for a deliberately conservative reading posture in the non-epileptic context: when a transient does not clearly meet the morphological criteria for an epileptiform discharge, it should be described neutrally (“sharply contoured transient of uncertain significance,” “consistent with wicket activity”) rather than labeled epileptiform.")
body("Two safeguards reduce the error rate. The first is the disciplined use of standardized terminology, which constrains readers to a shared, criterion-based vocabulary and resists the drift toward over-interpretation (Kane et al., 2017). The second is expert over-read: in any consequential case, and certainly before any clinical action is taken on an epileptiform finding discovered incidentally — including a finding surfaced by a QEEG or neurofeedback assessment — the record should be reviewed by an electroencephalographer experienced in epilepsy. The cost of a second opinion is trivial against the cost of a wrong diagnosis.")
divider()

# ============================== 7. CLINICAL MANAGEMENT ==============================
h1("7.  Clinical Management of Abnormal Paroxysmal Activity in Routine EEG")
body("When abnormal paroxysmal activity is encountered in a routine or screening EEG of a person without a seizure history, a small number of principles yield defensible management.")
body("*Diagnose from the clinic, not the tracing.* Epilepsy is not diagnosed from an EEG. The waveform is correlated with a careful history of paroxysmal clinical events; in their absence, an epileptiform finding is an isolated electrographic observation, and that is precisely how it should be recorded and conveyed.")
body("*Describe before you label.* The EEG report should state what was seen, where, in what state, and with what morphology, and should reserve the word “epileptiform” for transients that meet criteria. Equivocal transients are named as such. A line distinguishing an electrographic description from a clinical interpretation protects both patient and clinician.")
body("*Extend the question, not the conclusion.* When a finding is ambiguous or clinically important, the appropriate next step is more or better data — a repeat study, a sleep-deprived or sleep recording, or longer monitoring — not a firmer assertion from the same limited segment. If genuine epileptiform activity is confirmed, or if the clinical picture raises real concern, referral to an epileptologist is the correct path.")
body("*Counsel proportionately, and do not treat the EEG.* An asymptomatic person with an incidental discharge should be reassured in proportion to the modest predictive value such a finding carries, not alarmed. Antiseizure medication is not initiated for an incidental epileptiform discharge in the absence of clinical seizures; that decision belongs to a treating neurologist weighing the whole picture. Driving, occupational, and forensic implications should be handled conservatively and, where they arise, with specialist and where appropriate legal input — again recognizing that the gravest harm usually comes from over-calling, not under-calling, the finding.")
divider()

# ============================== 8. IMPLICATIONS FOR NFB ==============================
h1("8.  Implications for QEEG-Guided Neurofeedback")
body("These considerations are not academic for the neurofeedback practitioner; they are operational. A QEEG obtained to guide training is a resting-state EEG, and across a busy practice such recordings will inevitably surface sharply contoured transients — most of them benign variants, a few of them genuine. The first implication is procedural: the raw EEG behind any QEEG used for clinical decision-making should be inspected for epileptiform and abnormal paroxysmal activity by someone competent to recognize it, and any consequential finding should receive an expert over-read before training proceeds. A normative Z-score report is not a substitute for visual review, because the spectral pipeline that produces it is, by construction, insensitive to the brief epileptiform transient (Nuwer, 1997).")
body("The second implication concerns protocol selection and safety. A confirmed epileptiform finding does not necessarily contraindicate neurofeedback — indeed the companion chapter develops the evidence that EEG self-regulation can *raise* seizure threshold — but it does shift the choice of protocol toward the stabilizing end of the spectrum: sensorimotor-rhythm enhancement, suppression of excess slow activity, and Z-score normalization, rather than aggressive uptraining of fast activity, which is generally avoided in a brain that has shown a tendency to paroxysmal discharge. It also mandates coordination: where genuine epileptiform activity is present, training is conducted with the awareness and, ideally, the involvement of the treating physician, and with clear documentation of the finding and the plan.")
body("In this way the detection and interpretation of epileptiform activity in people without epilepsy is not a separate concern from neuromodulation but its necessary first step. The clinical reading of the EEG establishes who can be trained, with what protocol, and with what precautions — the foundation on which the non-invasive neuromodulation interventions of the companion chapter (Collura, Ims, & Turner, this volume) are built.")
divider()

# ============================== 9. CONCLUSION ==============================
h1("9.  Conclusion")
body("Epileptiform-appearing activity is part of the normal range of human EEG variation, not the exclusive signature of epilepsy. It is found, at low but real rates, in healthy adults and at higher rates in children and in several neurodevelopmental and psychiatric populations, and in the person without a seizure history its predictive value is modest. The central clinical skills are therefore two: the morphological discipline to separate genuine epileptiform discharges from the benign variants that mimic them, and the interpretive restraint to keep an isolated discharge from becoming an unwarranted diagnosis. Standardized terminology, conservative reporting, expert over-read, and management driven by the clinical history rather than the tracing are the safeguards that turn a potentially harmful incidental finding into a correctly weighted piece of information — and that prepare the EEG, and the patient, for the neuromodulation approaches that follow.")
divider()

# ============================== ACK / DISCLOSURE ==============================
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after = Pt(3)
run(p, "Acknowledgments", bold=True, color=C_TITLE, size=12.5)
body("The authors report no conflicts of interest relevant to the content of this chapter.")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
run(p, "AI-use disclosure: ", bold=True)
run(p, "Generative AI (Claude) was used only as an assistive drafting and formatting aid in preparing this chapter; all scientific content, clinical interpretation, citations, and conclusions are the authors’ own and were reviewed, revised, and independently verified by the authors.", italic=False)
divider()

# ============================== REFERENCES ==============================
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after = Pt(4)
run(p, "References", bold=True, color=C_TITLE, size=12.5)

refs = [
 "Benbadis, S. R. (2007). Errors in EEGs and the misdiagnosis of epilepsy: Importance, causes, consequences, and proposed remedies. Epilepsy & Behavior, 11(3), 257–262.",
 "Benbadis, S. R., & Tatum, W. O. (2003). Overinterpretation of EEGs and misdiagnosis of epilepsy. Journal of Clinical Neurophysiology, 20(1), 42–44.",
 "Cavazzuti, G. B., Cappella, L., & Nalin, A. (1980). Longitudinal study of epileptiform EEG patterns in normal children. Epilepsia, 21(1), 43–55.",
 "Fisher, R. S., Cross, J. H., French, J. A., Higurashi, N., Hirsch, E., Jansen, F. E., … Zuberi, S. M. (2017). Operational classification of seizure types by the International League Against Epilepsy. Epilepsia, 58(4), 522–530.",
 "Gregory, R. P., Oates, T., & Merry, R. T. G. (1993). Electroencephalogram epileptiform abnormalities in candidates for aircrew training. Electroencephalography and Clinical Neurophysiology, 86(1), 75–77.",
 "Kane, N., Acharya, J., Beniczky, S., Caboclo, L., Finnigan, S., Kaplan, P. W., … Pressler, R. (2017). A revised glossary of terms most commonly used by clinical electroencephalographers and updated proposal for the report format of the EEG findings. Clinical Neurophysiology Practice, 2, 170–185.",
 "Kwan, P., Arzimanoglou, A., Berg, A. T., Brodie, M. J., Allen Hauser, W., Mathern, G., … French, J. (2010). Definition of drug resistant epilepsy: Consensus proposal by the ad hoc Task Force of the ILAE Commission on Therapeutic Strategies. Epilepsia, 51(6), 1069–1077.",
 "Nuwer, M. (1997). Assessment of digital EEG, quantitative EEG, and EEG brain mapping: Report of the American Academy of Neurology and the American Clinical Neurophysiology Society. Neurology, 49(1), 277–292.",
 "Pascual-Marqui, R. D., Michel, C. M., & Lehmann, D. (1994). Low resolution electromagnetic tomography: A new method for localizing electrical activity in the brain. International Journal of Psychophysiology, 18(1), 49–65.",
 "Pillai, J., & Sperling, M. R. (2006). Interictal EEG and the diagnosis of epilepsy. Epilepsia, 47(Suppl. 1), 14–22.",
 "Sam, M. C., & So, E. L. (2001). Significance of epileptiform discharges in patients without epilepsy in the community. Epilepsia, 42(10), 1273–1278.",
 "Scheffer, I. E., Berkovic, S., Capovilla, G., Connolly, M. B., French, J., Guilhoto, L., … Zuberi, S. M. (2017). ILAE classification of the epilepsies: Position paper of the ILAE Commission for Classification and Terminology. Epilepsia, 58(4), 512–521.",
 "Shelley, B. P., Trimble, M. R., & Boutros, N. N. (2008). Electroencephalographic cerebral dysrhythmic abnormalities in the trinity of nonepileptic general population, neuropsychiatric, and neurobehavioral disorders. Journal of Neuropsychiatry and Clinical Neurosciences, 20(1), 7–22.",
 "Smith, S. J. M. (2005). EEG in the diagnosis, classification, and management of patients with epilepsy. Journal of Neurology, Neurosurgery & Psychiatry, 76(Suppl. 2), ii2–ii7.",
 "So, E. L. (2010). Interictal epileptiform discharges in persons without a history of seizures: What do they mean? Journal of Clinical Neurophysiology, 27(4), 229–238.",
 "Spence, S. J., & Schneider, M. T. (2009). The role of epilepsy and epileptiform EEGs in autism spectrum disorders. Pediatric Research, 65(6), 599–606.",
 "Tatum, W. O., Husain, A. M., Benbadis, S. R., & Kaplan, P. W. (2006). Normal adult EEG and patterns of uncertain significance. Journal of Clinical Neurophysiology, 23(3), 194–207.",
 "Thatcher, R. W. (1998). Normative EEG databases and EEG biofeedback. Journal of Neurotherapy, 2(4), 8–39.",
 "Zivin, L., & Ajmone Marsan, C. (1968). Incidence and prognostic significance of “epileptiform” activity in the EEG of non-epileptic subjects. Brain, 91(4), 751–778.",
]
for r in refs:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4); p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    run(p, r, size=10.5)

page_number_footer()
doc.save(OUT)
print("Saved:", OUT)
print("Paragraphs:", len(doc.paragraphs))
